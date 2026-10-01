"""「明细」页面的自动化测试：筛选、编辑弹窗、导出、删除。

运行方式（在项目根目录）：
    conda activate vibe_coding_learning
    QT_QPA_PLATFORM=offscreen python tests/test_expense_list.py

注意：测试使用临时数据库（草稿纸），绝不会碰真实账本 data/account.db。
"""

import shutil
import sys
import tempfile
from pathlib import Path

# 让脚本无论从哪里运行都能找到项目代码包 account_app
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from PySide6.QtCore import QDate
from PySide6.QtWidgets import QApplication
from openpyxl import load_workbook

from account_app import db
from account_app.ui import expense_list as el_mod
from account_app.ui.expense_dialog import ExpenseDialog
from account_app.ui.expense_list import ExpenseList
from account_app.ui.main_window import MainWindow


def _cat_id(top_name: str, sub_name: str) -> int:
    """按名称查二级分类 id。"""
    conn = db.get_connection()
    top = conn.execute(
        "SELECT id FROM categories WHERE name=? AND parent_id IS NULL", (top_name,)
    ).fetchone()
    sub = conn.execute(
        "SELECT id FROM categories WHERE name=? AND parent_id=?", (sub_name, top["id"])
    ).fetchone()
    conn.close()
    return sub["id"]


def main() -> None:
    app = QApplication([])
    tmp = Path(tempfile.mkdtemp())
    try:
        # ---- 切到临时数据库（草稿纸模式） ----
        db.DB_PATH = tmp / "test.db"
        el_mod.EXPORT_DIR = tmp / "exports"
        db.init_db()

        # ---- 准备测试数据：3 笔账 ----
        id_a = db.insert_expense(800, _cat_id("餐饮", "早餐"), "2026-10-01", "公司楼下")
        id_b = db.insert_expense(400, _cat_id("交通", "公交/地铁"), "2026-09-30", "上班")
        id_c = db.insert_expense(5000, _cat_id("购物", "日用品"), "2026-09-28", "超市采购")

        # ---- 筛选测试 ----
        page = ExpenseList()

        def check(desc: str, expect_ids: list[int]) -> None:
            got = sorted(r["id"] for r in page._last_rows)
            assert got == sorted(expect_ids), f"{desc}：期望 {expect_ids}，实际 {got}"
            print(f"  ✅ {desc} → {got}")

        page.keyword_input.setText("超市")
        page.refresh()
        check("关键字「超市」", [id_c])

        page.keyword_input.clear()
        page.date_check.setChecked(True)
        page.date_from.setDate(QDate(2026, 9, 30))
        page.date_to.setDate(QDate(2026, 10, 1))
        page.refresh()
        check("日期 09-30 ~ 10-01", [id_a, id_b])
        page.date_check.setChecked(False)

        page.top_filter.setCurrentIndex(page.top_filter.findText("交通"))
        page.refresh()
        check("分类筛选「交通」", [id_b])
        page.top_filter.setCurrentIndex(0)
        assert not page.sub_filter.isEnabled()
        print("  ✅ 大类=全部时小类自动禁用")

        page.min_amount.setValue(10.0)
        page.refresh()
        check("最小金额 ¥10", [id_c])
        page.min_amount.setValue(0.0)

        page.refresh()
        assert page.table.rowCount() == 3
        assert "共 3 笔" in page.summary.text() and "62.00" in page.summary.text()
        print(f"  ✅ 汇总行正常：{page.summary.text()}")

        # ---- 编辑弹窗测试 ----
        dlg = ExpenseDialog(id_a)
        assert dlg.amount_input.value() == 8.00
        assert dlg.top_cat.currentText() == "餐饮"
        assert dlg.sub_cat.currentText() == "早餐"
        assert dlg.note_input.text() == "公司楼下"
        print("  ✅ 编辑弹窗回填正确：¥8.00 餐饮·早餐「公司楼下」")
        dlg.amount_input.setValue(9.99)
        dlg.note_input.setText("改过备注")
        dlg._on_save()
        row = db.get_expense(id_a)
        assert row["amount_cents"] == 999 and row["note"] == "改过备注"
        print("  ✅ 编辑保存成功：改为 ¥9.99「改过备注」")

        # ---- 导出测试 ----
        page.refresh()
        page._export("csv")
        page._export("xlsx")
        files = sorted((tmp / "exports").glob("账目导出_*"))
        assert len(files) == 2, f"应导出 2 个文件，实际 {len(files)}"
        csv_file = next(f for f in files if f.suffix == ".csv")
        xlsx_file = next(f for f in files if f.suffix == ".xlsx")
        with open(csv_file, encoding="utf-8-sig") as f:
            lines = f.read().strip().splitlines()
        assert len(lines) == 4 and lines[0].startswith("日期")
        assert "超市采购" in lines[-1]
        wb = load_workbook(xlsx_file)
        assert wb.active.max_row == 4
        print("  ✅ 导出正常：CSV / Excel 各 4 行（含表头）")

        # ---- 删除 + 主窗口联动测试 ----
        db.delete_expense(id_b)
        win = MainWindow()
        assert win.expense_list.table.rowCount() == 2
        print("  ✅ 删除正常；主窗口与页面联动正常")

        print("全部测试通过 ✅")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
